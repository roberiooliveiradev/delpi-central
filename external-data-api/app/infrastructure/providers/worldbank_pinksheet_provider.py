"""World Bank "Pink Sheet" monthly commodity prices provider (PROC-0076).

Downloads the official monthly workbook
(``CMO-Historical-Data-Monthly.xlsx``) and extracts normalized monthly
observations for catalogued metal instruments. The World Bank licenses
its datasets under CC BY 4.0 — attribution rides in the source catalog.

Workbook layout (verified against the October-2026 file): sheet
"Monthly Prices", one header row with commodity display names
(``Copper``, ``Aluminum``, ``Zinc``, ``Nickel``, ``Lead``, ``Tin`` …),
the next row carries units (``($/mt)``), then one row per month keyed
``YYYYMmm`` (e.g. ``1960M01``) in column A. Missing values appear as
``…`` or blank cells and are skipped — never ingested as zero.
"""

from __future__ import annotations

import io
import logging
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any, Callable, Iterable

import requests

from app.application.contracts import (
    InvalidResponseError,
    ProviderTimeoutError,
    ProviderUnavailableError,
    RateLimitedError,
)
from app.domain.market_data import PROVIDER_WORLDBANK_PINK_SHEET
from app.infrastructure.config.settings import Settings

logger = logging.getLogger("external-data-api.pinksheet")

Transport = Callable[..., Any]

_SHEET_NAME = "Monthly Prices"
_PERIOD_KEY = re.compile(r"^(\d{4})M(\d{2})$")
_UNIT_EXPECTED = re.compile(r"^\(?\s*\$/mt\s*\)?$", re.IGNORECASE)
# Max rows scanned when locating the header row defensively.
_HEADER_SCAN_LIMIT = 15


def _requests_transport(url: str, *, headers: dict, timeout: float) -> Any:
    return requests.get(url, headers=headers, timeout=timeout)


class WorldBankPinkSheetProvider:
    """Monthly metal observations from the Pink Sheet workbook."""

    provider_id = PROVIDER_WORLDBANK_PINK_SHEET

    def __init__(
        self,
        *,
        workbook_url: str | None = None,
        timeout_seconds: float | None = None,
        transport: Transport | None = None,
    ) -> None:
        self._url = (
            workbook_url or Settings.WORLDBANK_PINKSHEET_MONTHLY_URL
        )
        self._timeout = (
            timeout_seconds
            if timeout_seconds is not None
            else Settings.WORLDBANK_PINKSHEET_TIMEOUT_SECONDS
        )
        self._transport = transport or _requests_transport

    def fetch_period(
        self,
        instruments: Iterable[Any],
        date_from: date,
        date_to: date,
    ) -> list[dict[str, Any]]:
        """Download + normalize observations for the requested instruments.

        ``instruments`` carry ``source_series_key`` (workbook column name)
        and ``instrument_code``. Returns rows shaped for upsert:
        ``instrument_code, observation_date, price, source_record_key``.
        """
        wanted = {
            inst.source_series_key: inst.instrument_code
            for inst in instruments
        }
        if not wanted:
            return []

        body = self._download()
        return self._normalize(body, wanted, date_from, date_to)

    # --- internals ---------------------------------------------------------

    def _download(self) -> bytes:
        headers = {
            "Accept": (
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet, application/octet-stream"
            ),
            "User-Agent": Settings.EXTERNAL_DATA_USER_AGENT,
        }
        try:
            response = self._transport(
                self._url, headers=headers, timeout=self._timeout
            )
        except requests.Timeout as exc:
            raise ProviderTimeoutError("pink sheet download timed out") from exc
        except requests.RequestException as exc:
            logger.warning("pinksheet_transport_failure")
            raise ProviderUnavailableError("pink sheet download failed") from exc

        status = getattr(response, "status_code", 0)
        if status == 429:
            retry_after = None
            try:
                retry_after = int(response.headers.get("Retry-After") or 0)
            except (TypeError, ValueError):
                retry_after = None
            raise RateLimitedError(
                "pink sheet rate limited", retry_after=retry_after or None
            )
        if status in (401, 403):
            raise ProviderUnavailableError("pink sheet access denied")
        if status >= 500:
            raise ProviderUnavailableError("pink sheet server error")
        if status != 200:
            raise ProviderUnavailableError(
                f"pink sheet unexpected status {status}"
            )

        content = getattr(response, "content", b"")
        if not content:
            raise InvalidResponseError("pink sheet empty body")
        return bytes(content)

    def _normalize(
        self,
        body: bytes,
        wanted: dict[str, str],
        date_from: date,
        date_to: date,
    ) -> list[dict[str, Any]]:
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover — dependency guard
            raise ProviderUnavailableError("openpyxl not installed") from exc

        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(body), read_only=True, data_only=True
            )
        except Exception as exc:
            raise InvalidResponseError("pink sheet unreadable workbook") from exc

        if _SHEET_NAME not in workbook.sheetnames:
            raise InvalidResponseError(
                f"pink sheet sheet '{_SHEET_NAME}' not found"
            )
        sheet = workbook[_SHEET_NAME]

        rows = sheet.iter_rows(values_only=True)
        header_row, units_row = self._find_header(rows, wanted)

        # Map wanted series names to column indexes.
        col_by_key: dict[str, int] = {}
        for idx, cell in enumerate(header_row):
            name = str(cell).strip() if cell is not None else ""
            if name in wanted:
                col_by_key[name] = idx
                unit = (
                    str(units_row[idx]).strip()
                    if idx < len(units_row) and units_row[idx] is not None
                    else ""
                )
                if not _UNIT_EXPECTED.match(unit):
                    raise InvalidResponseError(
                        f"pink sheet unexpected unit for {name}: {unit!r}"
                    )
        missing = [key for key in wanted if key not in col_by_key]
        if missing:
            raise InvalidResponseError(
                f"pink sheet columns not found: {sorted(missing)}"
            )

        observations: list[dict[str, Any]] = []
        for row in rows:
            period_cell = row[0] if row else None
            match = _PERIOD_KEY.match(
                str(period_cell).strip() if period_cell is not None else ""
            )
            if not match:
                continue
            year, month = int(match.group(1)), int(match.group(2))
            obs_date = date(year, month, 1)
            if obs_date < date_from or obs_date > date_to:
                continue
            for key, idx in col_by_key.items():
                raw = row[idx] if idx < len(row) else None
                price = self._parse_price(raw)
                if price is None:
                    # Legitimate gap — skip, never fabricate.
                    continue
                observations.append(
                    {
                        "instrument_code": wanted[key],
                        "observation_date": obs_date,
                        "price": price,
                        "source_record_key": f"pink-sheet:{year}M{month:02d}:{key}",
                    }
                )
        return observations

    def _find_header(self, rows, wanted: dict[str, str]):
        """Locate (header, units) rows defensively near the top."""
        seen: list[tuple] = []
        for _ in range(_HEADER_SCAN_LIMIT):
            try:
                row = next(rows)
            except StopIteration:
                break
            seen.append(row)
            names = {
                str(c).strip() for c in row if isinstance(c, str)
            }
            if any(key in names for key in wanted):
                try:
                    units_row = next(rows)
                except StopIteration as exc:
                    raise InvalidResponseError(
                        "pink sheet truncated after header"
                    ) from exc
                return row, units_row
        raise InvalidResponseError("pink sheet header row not found")

    @staticmethod
    def _parse_price(raw: Any) -> Decimal | None:
        if raw is None:
            return None
        if isinstance(raw, str):
            text = raw.strip()
            if not text or text in {"…", "..", "—", "-", "n/a"}:
                return None
            try:
                value = Decimal(text)
            except InvalidOperation:
                return None
        else:
            try:
                value = Decimal(str(raw))
            except InvalidOperation:
                return None
        return value if value > 0 else None


def build_pinksheet_provider(**overrides) -> WorldBankPinkSheetProvider:
    return WorldBankPinkSheetProvider(**overrides)
